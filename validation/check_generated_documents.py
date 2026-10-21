"""Open the documents actually produced by Android using independent Office/PDF readers."""
import csv
import json
from pathlib import Path
import sys
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

folder = Path(sys.argv[1])
word = Document(folder / "document.docx")
assert "Bonjour é & <Android>" in "\n".join(p.text for p in word.paragraphs)
book = load_workbook(folder / "document.xlsx")
sheet = book.active
assert sheet["A2"].value == "é & <Android>"
assert sheet["B2"].value == 42
assert sheet["B3"].value is True
assert sheet["A3"].value == "=1+1" and sheet["A3"].data_type == "s"
slides = Presentation(folder / "document.pptx")
assert len(slides.slides) == 2
text = "\n".join(shape.text for slide in slides.slides for shape in slide.shapes if shape.has_text_frame)
assert "Bonjour é & <Android>" in text and "Deuxième paragraphe" in text
pdf = PdfReader(folder / "document.pdf")
assert len(pdf.pages) == 1 and "Bonjour" in pdf.pages[0].extract_text()
assert json.loads((folder / "document.json").read_text())["text"] == "Bonjour é & <Android>"
with (folder / "document.csv").open(newline="") as stream:
    assert list(csv.reader(stream))[1] == ["é & <Android>", "42"]
assert "Bonjour é & <Android>" in (folder / "document.txt").read_text()
print("Android-generated PDF, DOCX, XLSX, PPTX, JSON, CSV and TXT opened successfully by independent readers.")
